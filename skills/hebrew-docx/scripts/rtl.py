#!/usr/bin/env python3
"""Mixed Hebrew/English paragraph helpers for python-docx, verified against Microsoft Word.

These rules were found by rendering isolated variants side by side in Word and checking them
by eye, not derived from UAX#9 on paper. They deliberately differ from the commonly published
rule "never put <w:rtl/> on runs of a mixed Hebrew+English paragraph". See ../SKILL.md for the
full reasoning and the phrasing patterns that still need rewording.

Pipeline for one paragraph of text:
  1. _split_atomic_mixed_brackets: a "(...)" group whose inner text mixes a Hebrew word and a
     Latin word is emitted as ONE run with no rtl flag.
  2. everything else: _split_by_script -> _shift_boundary_spaces -> _shift_boundary_brackets,
     and every Hebrew segment gets <w:rtl/> unconditionally.
  3. every run gets w:cs / w:szCs (and w:bCs for bold), so font and size apply to Hebrew.

Usage:
    from docx import Document
    from rtl import add_para
    doc = Document()
    add_para(doc, 'אנליסטית נתונים | SQL, Power BI ואוטומציה', size=10, bold=True)
"""
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

_HEB = re.compile(r'[֐-׿יִ-ﭏ]')
_OPEN_BR, _CLOSE_BR = '([{/-', ')]}/-'
_MIXED_BRACKET = re.compile(r'\([^()]*\)')


def has_hebrew(text):
    return bool(_HEB.search(text))


def _split_atomic_mixed_brackets(text):
    """Split text into (chunk, is_atomic_mixed_bracket) pieces.

    A "(...)" span whose INNER content has both a Hebrew letter and a Latin letter
    (e.g. "(דוחות Power BI)") must be emitted as ONE run with no w:rtl flag: splitting it by
    script, or flagging it rtl, makes Word scramble the word order inside it or mirror the
    bracket glyphs. A bracket that is purely Hebrew inside (e.g. "(משרה חלקית)") or purely
    Latin/digits inside (e.g. "(SAP)", "(כ-6 חודשים)" - digits are not Latin letters) is NOT
    atomic and goes through the normal per-script pipeline, which handles it correctly.
    """
    pieces = []
    last = 0
    for m in _MIXED_BRACKET.finditer(text):
        inner = m.group(0)[1:-1]
        has_heb = bool(_HEB.search(inner))
        has_lat = any(ch.isascii() and ch.isalpha() for ch in inner)
        if has_heb and has_lat:
            if m.start() > last:
                pieces.append((text[last:m.start()], False))
            pieces.append((m.group(0), True))
            last = m.end()
    if last < len(text):
        pieces.append((text[last:], False))
    return pieces or [(text, False)]


def _strong(ch):
    """True for a Hebrew letter, False for a strong-LTR char (Latin letter or ASCII digit),
    None for a neutral (space, punctuation). Digits count as LTR so a number never rides
    inside a Hebrew run."""
    if _HEB.match(ch):
        return True
    if ch.isascii() and ch.isalnum():
        return False
    return None


def _split_by_script(text):
    """Split a string into (segment, is_rtl) runs. Neutrals attach to the current run;
    leading neutrals take the direction of the first strong character in the whole string."""
    default_rtl = next((s for s in (_strong(c) for c in text) if s is not None), True)
    segments, buf, buf_rtl = [], '', None
    for ch in text:
        s = _strong(ch)
        kind = s if s is not None else (buf_rtl if buf_rtl is not None else default_rtl)
        if buf_rtl is None or kind == buf_rtl:
            buf, buf_rtl = buf + ch, kind
        else:
            segments.append((buf, buf_rtl))
            buf, buf_rtl = ch, kind
    if buf:
        segments.append((buf, buf_rtl))
    return segments


_TRAILING_NEUTRALS = ' |–-·:/'


def _shift_boundary_spaces(segments):
    """Move the trailing neutrals (space, pipe, dashes, middle dot, colon, slash) at the END of an
    LTR run that precedes an RTL run to the START of that RTL run.

    Spaces: Word trims trailing whitespace at a direction boundary, gluing words together.
    Separators: left in the LTR run, Word draws them on the wrong side of the line. Verified in
    Word (cv-tailor diagnostic renders): "דצמבר 2024 – היום" had the dash before 2024,
    "Example Data Academy  |  מאי 2024" and "050-0000000  |  x@y.com  |  תל אביב" had the pipe
    at the line edge; all three render correctly with the separator moved, and lines that were
    already correct (title lines, skills, bullets, brackets, multi-clause) did not change.
    A trailing comma is NOT moved: that pattern still breaks, and is fixed by rewording.
    """
    out = [[seg, rtl] for seg, rtl in segments]
    for i in range(len(out) - 1):
        seg, rtl = out[i]
        nseg, nrtl = out[i + 1]
        if rtl is False and nrtl is True and seg.rstrip(_TRAILING_NEUTRALS) not in ('', seg):
            stripped = seg.rstrip(_TRAILING_NEUTRALS)
            out[i][0] = stripped
            out[i + 1][0] = seg[len(stripped):] + nseg
    return [(s, r) for s, r in out if s]


def _shift_boundary_brackets(segments):
    """Move an opening bracket at the end of a run into the next run, and a closing bracket at
    the start of a run into the previous run, whenever the two runs differ in direction.

    Must be symmetric: both "Hebrew ends with '(' -> English" (e.g. "נתונים (Big Data)") and
    its mirror "English/number ends with '(' -> Hebrew" (e.g. "3 (כ-6 חודשים)") need the
    move. Handling only one direction leaves the other with mirrored brackets in Word.
    """
    out = [[seg, rtl] for seg, rtl in segments]
    for i in range(len(out) - 1):
        seg, rtl = out[i]
        nseg, nrtl = out[i + 1]
        stripped = seg.rstrip(' ')
        if rtl != nrtl and stripped and stripped[-1] in _OPEN_BR:
            out[i][0] = stripped[:-1]
            out[i + 1][0] = seg[len(stripped) - 1:] + nseg
    for i in range(len(out) - 1, 0, -1):
        seg, rtl = out[i]
        pseg, prtl = out[i - 1]
        lstripped = seg.lstrip(' ')
        if rtl != prtl and lstripped and lstripped[0] in _CLOSE_BR:
            lead = len(seg) - len(lstripped)
            out[i][0] = lstripped[1:]
            out[i - 1][0] = pseg + seg[:lead + 1]
    return [(s, r) for s, r in out if s]


def _para_is_rtl(text):
    """Any Hebrew -> RTL paragraph. Latin but no Hebrew -> LTR. All-neutral -> RTL."""
    if _HEB.search(text):
        return True
    if any(ch.isascii() and ch.isalnum() for ch in text):
        return False
    return True


def rtl_segments(text):
    """Return the final list of (run_text, rtl_flag) for a paragraph, after every rule."""
    result = []
    for chunk, is_atomic in _split_atomic_mixed_brackets(text):
        if is_atomic:
            result.append((chunk, False))
        else:
            result.extend(_shift_boundary_brackets(_shift_boundary_spaces(_split_by_script(chunk))))
    return result


def _emit_run(p, text, rtl_flag, font, size, bold, color):
    run = p.add_run(text)
    rPr = run._r.get_or_add_rPr()
    # rPr children in OOXML schema order: rFonts, b, bCs, color, sz, szCs, rtl
    rPr.append(rPr.makeelement(qn('w:rFonts'), {
        qn('w:ascii'): font, qn('w:hAnsi'): font, qn('w:cs'): font}))
    if bold:
        rPr.append(rPr.makeelement(qn('w:b'), {}))
        rPr.append(rPr.makeelement(qn('w:bCs'), {}))
    if color:
        rPr.append(rPr.makeelement(qn('w:color'), {qn('w:val'): color}))
    half_pt = str(int(round(size * 2)))
    rPr.append(rPr.makeelement(qn('w:sz'), {qn('w:val'): half_pt}))
    rPr.append(rPr.makeelement(qn('w:szCs'), {qn('w:val'): half_pt}))
    if rtl_flag:
        rPr.append(rPr.makeelement(qn('w:rtl'), {}))
    return run


def add_para(container, text, font='Arial', size=10, bold=False, color=None, after=0,
             before=0, border_bottom=None, align=None, force_rtl=False, indent=None):
    """Add one paragraph of possibly mixed Hebrew/English text.

    text: a string, or a list of (text, bold, color) parts for mixed styling in one paragraph.
      Direction is decided on the WHOLE line (the parts joined), then each run is cut at the
      part boundaries, so styling never changes how Word orders the line.
    align: None (natural for the direction), 'center', or 'left' (forces an LTR paragraph).
    border_bottom: a hex color for a bottom rule under the paragraph (section headers).
    indent: twips for a hanging indent (bullets), measured from the paragraph's start side.
    """
    parts = [(text, bold, color)] if isinstance(text, str) else list(text)
    text = ''.join(part[0] for part in parts)
    p = container.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    # pPr children in schema order: pBdr, bidi, spacing, ind (jc is placed by python-docx)
    if border_bottom:
        pBdr = pPr.makeelement(qn('w:pBdr'), {})
        pBdr.append(pBdr.makeelement(qn('w:bottom'), {
            qn('w:val'): 'single', qn('w:sz'): '6', qn('w:space'): '2',
            qn('w:color'): border_bottom}))
        pPr.append(pBdr)
    base_rtl = force_rtl or (_para_is_rtl(text) and align != 'left')
    if base_rtl:
        pPr.append(pPr.makeelement(qn('w:bidi'), {}))
    pPr.append(pPr.makeelement(qn('w:spacing'), {
        qn('w:before'): str(int(before)), qn('w:after'): str(int(after))}))
    if indent:
        # In a <w:bidi/> paragraph Word applies w:left to the start (right) side.
        pPr.append(pPr.makeelement(qn('w:ind'), {
            qn('w:left'): str(indent), qn('w:hanging'): str(indent)}))
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'left':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # part boundaries as character offsets into the joined text
    bounds, pos = [], 0
    for part_text, part_bold, part_color in parts:
        bounds.append((pos, pos + len(part_text), part_bold, part_color))
        pos += len(part_text)
    pos = 0
    for seg, is_rtl in rtl_segments(text):
        seg_start, seg_end = pos, pos + len(seg)
        for b_start, b_end, part_bold, part_color in bounds:
            lo, hi = max(seg_start, b_start), min(seg_end, b_end)
            if lo < hi:
                _emit_run(p, text[lo:hi], is_rtl, font, size, part_bold, part_color)
        pos = seg_end
    return p
